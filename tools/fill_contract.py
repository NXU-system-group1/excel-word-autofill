# -*- coding: utf-8 -*-
"""填充「6、合同」：只替换红字信息里有的乙方字段，甲方一律不动。

合同段10 的占位写成「联系人」但实际是法定代表人栏，段11 是「联系人及联系电话」，
这两处按段落上下文纠正，避免填错人。
"""
import sys
from pathlib import Path

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from doc2docx import convert as doc_to_docx  # noqa: E402

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:/git学习/Excel转指定位置word")
BASE = ROOT / "项目文件"
INFO = BASE / "红字信息" / "晨早.xlsx"
SRC_DOC = BASE / "1、工程A三方询价 价格" / "6、合同.doc"
TMP_DOCX = ROOT / ".tmp" / "合同.docx"
OUT_DOCX = BASE / "1、工程A三方询价 价格" / "6、合同_已填.docx"


def is_red(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    c = rpr.find(qn("w:color"))
    if c is None:
        return False
    val = c.get(qn("w:val"))
    if not val or len(val) != 6:
        return False
    try:
        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except ValueError:
        return False
    return r >= 0x80 and g <= 0x60 and b <= 0x60


def red_blocks(para):
    out, buf = [], []
    for r in para.runs:
        if is_red(r):
            buf.append(r)
        elif buf:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                k = cell._tc.getroottree().getpath(cell._tc)
                if k in seen:
                    continue
                seen.add(k)
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def fmt(val, field=""):
    if isinstance(val, (int, float)) and "金额" in str(field) and "大写" not in str(field):
        return f"{float(val):.2f}"
    if isinstance(val, float):
        return f"{val:.2f}"
    return str(val).strip()


def load_info(xlsx_path):
    """返回 (按占位取值, 按字段名取值)。"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    by_ph, by_field = {}, {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 3:
            continue
        ph, field, val = row[0], row[1], row[2]
        if val is None:
            continue
        text = fmt(val, field)
        if ph and str(ph).strip() not in ("—", ""):
            by_ph[str(ph).strip()] = text
        if field:
            by_field[str(field).strip()] = text
    wb.close()
    return by_ph, by_field


def resolve(para_text, red_text, by_ph, by_field):
    """按段落上下文纠正模板占位错误。"""
    # 段10「法定代表人：联系人」—— 该格要填法定代表人
    if "法定代表人" in para_text and red_text == "联系人":
        return by_field.get("乙方法定代表人")
    # 段11「联系人及联系电话：联系电话」—— 该格要填「姓名 电话」
    if "联系人及联系电话" in para_text and red_text == "联系电话":
        name = by_field.get("乙方联系人", "")
        tel = by_field.get("乙方电话", "")
        combo = f"{name} {tel}".strip()
        return combo or None
    return by_ph.get(red_text)


by_ph, by_field = load_info(INFO)
print(f"载入 晨早.xlsx：占位 {len(by_ph)} 条 / 字段 {len(by_field)} 条")

if not TMP_DOCX.exists() or TMP_DOCX.stat().st_mtime < SRC_DOC.stat().st_mtime:
    print("转换 .doc → .docx ...")
    doc_to_docx(SRC_DOC, TMP_DOCX)

doc = Document(str(TMP_DOCX))
hits, miss = [], []
for para in all_paragraphs(doc):
    ptext = para.text
    for buf in red_blocks(para):
        text = "".join(r.text for r in buf).strip()
        if not text:
            continue
        new = resolve(ptext, text, by_ph, by_field)
        if new is None:
            miss.append(text)
            continue
        buf[0].text = new
        buf[0].font.color.rgb = RGBColor(0, 0, 0)
        for r in buf[1:]:
            r.text = ""
        hits.append((text, new, ptext[:24]))

OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT_DOCX))

print("\n" + "=" * 74)
print(f"已替换 {len(hits)} 处：")
for old, new, ctx in hits:
    print(f"  {old}  →  {new}      〔{ctx}〕")

if miss:
    print(f"\n未替换 {len(miss)} 处（红字信息中无此项，保持红字）：")
    from collections import Counter
    for t, n in Counter(miss).most_common():
        label = t if len(t) <= 30 else t[:30] + "…"
        print(f"  {label}  ×{n}")

print("\n输出：", OUT_DOCX)
