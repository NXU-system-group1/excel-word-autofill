# -*- coding: utf-8 -*-
"""把报价清单的明细表同步进 Word 文档，行数随报价单自动增减。

用法：
    python tools/sync_quote_table.py <报价单.xlsx> <目标1.docx> [目标2.docx ...]

行为：
  - 从报价单里自动找表头行、明细行（序号为数字的行）、合计行，行数不限
  - 目标文档里定位表头与合计行，按报价单行数增删表格行（复制现有行格式，不重建）
  - 逐格写入 序号/名称/规格/数量/单位/单价/总价，文字一律改黑
  - 合计行的大写与小写按报价单重算

注意：只改目标文档的表格内容，不动其他段落；源 .doc 请先用 doc2docx 转好。
"""
import copy
import re
import sys
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import RGBColor

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

CN_NUM = "零壹贰叁肆伍陆柒捌玖"
CN_UNIT = ["", "拾", "佰", "仟"]
CN_BIG = ["", "万", "亿"]


def rmb_upper(amount) -> str:
    """金额转人民币大写。"""
    n = round(float(amount) + 1e-9, 2)
    yuan = int(n)
    cents = int(round((n - yuan) * 100))
    jiao, fen = divmod(cents, 10)

    if yuan == 0:
        int_part = "零"
    else:
        groups = []
        x = yuan
        while x > 0:
            groups.append(x % 10000)
            x //= 10000
        int_part = ""
        for idx in range(len(groups) - 1, -1, -1):
            g = groups[idx]
            if g == 0:
                if idx > 0 and int_part and not int_part.endswith("零"):
                    int_part += "零"
                continue
            seg, zero = "", False
            for pos in range(3, -1, -1):
                d = (g // 10 ** pos) % 10
                if d == 0:
                    if seg:
                        zero = True
                else:
                    if zero:
                        seg += "零"
                        zero = False
                    seg += CN_NUM[d] + CN_UNIT[pos]
            if idx < len(groups) - 1 and g < 1000 and not int_part.endswith("零"):
                int_part += "零"
            int_part += seg + CN_BIG[idx]

    if jiao == 0 and fen == 0:
        return int_part + "元整"
    res = int_part + "元"
    if jiao:
        res += CN_NUM[jiao] + "角"
    elif fen:
        res += "零"
    if fen:
        res += CN_NUM[fen] + "分"
    else:
        res += "整"
    return res


def read_quote(xlsx_path):
    """返回 (明细行列表, 合计金额, 合计大写)。

    明细行固定 7 项：序号 / 名称 / 参数 / 数量 / 单位 / 单价 / 总价。
    **按表头文字定位列**，不假设列顺序、不要求任何颜色标记——上传方给的
    Excel 只要第 1 列是「序号」、有「合计」行即可。
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        if str(ws.cell(r, 1).value or "").strip() == "序号":
            header_row = r
            break
    if header_row is None:
        wb.close()
        raise ValueError(f"{xlsx_path.name}：没找到表头行（第 1 列应为「序号」）")

    header = {}
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(header_row, c).value or "").strip()
        if v:
            header[c] = v

    def col(*keys, default=None):
        for c, name in header.items():
            for k in keys:
                if k in name:
                    return c
        return default

    c_no = col("序号", default=1)
    c_name = col("产品名称", "名称", default=2)
    c_spec = col("参数", "技术要求", "服务要求", "规格", default=3)
    c_qty = col("数量")
    c_unit = col("单位")
    c_price = col("单价")
    c_amt = col("总价", "金额")

    rows, total_amt, total_cn, total_row = [], None, None, None
    for r in range(header_row + 1, ws.max_row + 1):
        first = str(ws.cell(r, c_no).value or "").strip()
        joined = re.sub(r"\s+", "", "".join(str(ws.cell(r, c).value or "") for c in header))
        if "合计" in joined:
            total_row = r
            break
        if first.isdigit():
            vals = [
                ws.cell(r, c_no).value,
                ws.cell(r, c_name).value,
                ws.cell(r, c_spec).value if c_spec else None,
                ws.cell(r, c_qty).value if c_qty else None,
                ws.cell(r, c_unit).value if c_unit else None,
                ws.cell(r, c_price).value if c_price else None,
                ws.cell(r, c_amt).value if c_amt else None,
            ]
            # 总价常是公式（=D*F），缓存值可能缺失，用 数量×单价 兜底
            if vals[6] is None and vals[3] is not None and vals[5] is not None:
                try:
                    vals[6] = float(vals[3]) * float(vals[5])
                except (TypeError, ValueError):
                    pass
            rows.append(vals)

    if total_row:
        total_amt = ws.cell(total_row, c_amt).value if c_amt else None
        cn_text = ""
        for c in header:
            t = str(ws.cell(total_row, c).value or "")
            if "大写" in t:
                cn_text = t
                break
        m = re.search(r"大写[：:]\s*(.+)$", cn_text)
        total_cn = m.group(1).strip() if m else None
    wb.close()

    if not rows:
        raise ValueError(f"{xlsx_path.name}：表头下面没读到明细行")

    if total_amt is None:
        total_amt = sum(float(v[6] or 0) for v in rows)
    if not total_cn:
        total_cn = rmb_upper(total_amt)

    return rows, float(total_amt), total_cn


def uniq_cells(row):
    out, seen = [], set()
    for c in row.cells:
        k = c._tc.getroottree().getpath(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def row_texts(row):
    return [c.text.strip() for c in uniq_cells(row)]


def find_header(table):
    for i, row in enumerate(table.rows):
        t = row_texts(row)
        if t and t[0] == "序号":
            return i
    return None


def find_total(table, after):
    for i in range(after + 1, len(table.rows)):
        for t in row_texts(table.rows[i]):
            if "合计" in re.sub(r"\s+", "", t):
                return i
    return None


def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for r in p.runs[1:]:
            r.text = ""
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0, 0, 0)


def fmt_val(v, money=False):
    if v is None:
        return ""
    if money:
        try:
            return f"{float(v):.2f}"
        except (TypeError, ValueError):
            return str(v).strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def resize(table, header_idx, total_idx, need):
    """把数据区行数调整到 need 行，复制/删除表格行（保留格式）。返回新的合计行索引。"""
    cur = total_idx - header_idx - 1
    if cur == need:
        return total_idx

    if cur < need:
        for _ in range(need - cur):
            src = table.rows[total_idx - 1]._tr
            new = copy.deepcopy(src)
            src.addnext(new)
            total_idx += 1
    else:
        for _ in range(cur - need):
            tr = table.rows[total_idx - 1]._tr
            tr.getparent().remove(tr)
            total_idx -= 1
    return total_idx


def update_total(table, total_idx, amt, cn):
    """合计行：中文大写按正则替换；纯数字格里，只更新「等于原金额」的（合并单元格会出现多个）。"""
    cells = uniq_cells(table.rows[total_idx])

    nums = []
    for c in cells:
        t = c.text.strip()
        if re.fullmatch(r"[\d,]+(\.\d+)?", t):
            nums.append((c, float(t.replace(",", ""))))
    old = nums[0][1] if nums else None

    for c, v in nums:
        if old is not None and abs(v - old) < 0.01:
            set_cell(c, f"{amt:.2f}")

    for c in cells:
        text = c.text
        if not text.strip() or re.fullmatch(r"[\d,]+(\.\d+)?", text.strip()):
            continue
        new = re.sub(r"[零壹贰叁肆伍陆柒捌玖拾佰仟万亿元角分整]{3,}", cn, text)
        new = re.sub(
            r"(小写[：:]\s*[（(]?\s*[¥￥]?[：:]?\s*)[\d,]+(\.\d+)?",
            lambda m: m.group(1) + f"{amt:.2f}",
            new,
        )
        if new != text:
            set_cell(c, new)


def sync_one(docx_path, rows, amt, cn):
    doc = Document(str(docx_path))
    if not doc.tables:
        print(f"  [跳过] {docx_path.name}：没有表格")
        return False

    table = doc.tables[0]
    header_idx = find_header(table)
    if header_idx is None:
        print(f"  [跳过] {docx_path.name}：没找到「序号」表头")
        return False

    total_idx = find_total(table, header_idx)
    if total_idx is None:
        print(f"  [跳过] {docx_path.name}：没找到合计行")
        return False

    before = total_idx - header_idx - 1
    total_idx = resize(table, header_idx, total_idx, len(rows))

    for k, vals in enumerate(rows):
        cells = uniq_cells(table.rows[header_idx + 1 + k])
        for ci in range(min(7, len(cells))):
            set_cell(cells[ci], fmt_val(vals[ci], ci in (5, 6)))

    update_total(table, total_idx, amt, cn)
    doc.save(str(docx_path))
    print(f"  [完成] {docx_path.name}：数据行 {before} → {len(rows)}，合计 {amt:.2f}（{cn}）")
    return True


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)

    quote = Path(sys.argv[1])
    targets = [Path(p) for p in sys.argv[2:]]

    rows, amt, cn = read_quote(quote)
    print(f"报价单：{quote.name}")
    print(f"  明细 {len(rows)} 行，合计 {amt:.2f}，大写 {cn}")
    for r in rows:
        print("   ", " | ".join(fmt_val(v, i in (5, 6)) for i, v in enumerate(r)))

    print("\n写入目标：")
    ok = 0
    for t in targets:
        if not t.exists():
            print(f"  [跳过] {t.name}：文件不存在")
            continue
        if sync_one(t, rows, amt, cn):
            ok += 1
    print(f"\n完成 {ok}/{len(targets)} 份")


if __name__ == "__main__":
    main()
