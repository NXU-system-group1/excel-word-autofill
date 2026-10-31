# -*- coding: utf-8 -*-
"""按「红色系」扫描 docx 中的标记文本：R>=0x80 且 G<=0x60 且 B<=0x60。"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")


def is_red(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    c = rpr.find(qn("w:color"))
    if c is None:
        return False
    val = c.get(qn("w:val"))
    if not val or len(val) != 6:
        return False
    try:
        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except ValueError:
        return False
    return r >= 0x80 and g <= 0x60 and b <= 0x60


def color_of(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    c = rpr.find(qn("w:color"))
    return c.get(qn("w:val")) if c is not None else None


def blocks(para):
    out, buf = [], []
    for r in para.runs:
        if is_red(r):
            buf.append(r)
        elif buf:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out


def scan(path):
    doc = Document(str(path))
    items = []
    for i, p in enumerate(doc.paragraphs):
        for b in blocks(p):
            t = "".join(r.text for r in b).strip()
            if t:
                rest = p.text
                for r in b:
                    rest = rest.replace(r.text, "", 1)
                items.append((f"正文段{i}", rest.strip(" ：:，,。")[:40] or "(无上下文)",
                              t, color_of(b[0])))
    for ti, table in enumerate(doc.tables, 1):
        seen = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                k = cell._tc.getroottree().getpath(cell._tc)
                if k in seen:
                    continue
                seen.add(k)
                for p in cell.paragraphs:
                    for b in blocks(p):
                        t = "".join(r.text for r in b).strip()
                        if t:
                            label = ""
                            for lc in range(ci - 1, -1, -1):
                                s = row.cells[lc].text.strip()
                                if s:
                                    label = s.replace("\n", " ")
                                    break
                            items.append((f"表{ti} R{ri}C{ci}", label[:40] or "(无标签)",
                                          t, color_of(b[0])))
    for si, s in enumerate(doc.sections, 1):
        for nm, part in (("页眉", s.header), ("页脚", s.footer)):
            for p in part.paragraphs:
                for b in blocks(p):
                    t = "".join(r.text for r in b).strip()
                    if t:
                        items.append((f"{nm}{si}", "(页眉页脚)", t, color_of(b[0])))
    return items


for arg in sys.argv[1:]:
    p = Path(arg)
    print("\n" + "=" * 74)
    print("文件：", p.name)
    print("=" * 74)
    items = scan(p)
    if not items:
        print("  （无红色系标记）")
        continue
    for n, (w, lb, t, col) in enumerate(items, 1):
        print(f"{n:>2}. [{w:<12}] {col} <{lb}>  →  <{t}>")
    print(f"  合计 {len(items)} 处")
