# -*- coding: utf-8 -*-
"""公共层：路径常量、红色判定、文档遍历、红字映射、替换引擎。

所有 fill_* 脚本和 start.py 都从这里取，避免每个脚本各写一份。
"""
import re
from pathlib import Path

import openpyxl
from docx.oxml.ns import qn
from docx.shared import RGBColor

ROOT = Path(r"D:/git学习/Excel转指定位置word")
BASE = ROOT / "项目文件"
INFO_DIR = BASE / "红字信息"
DOC_DIR = BASE / "1、工程A三方询价 价格"
TMP_DIR = ROOT / ".tmp"

COMPANIES = ["晨早", "泉旭", "左右"]
BIDDER = "晨早"          # 中标方：晨早 44650
QUOTE_XLSX = DOC_DIR / "3.报价" / "1、报价清单（晨早）44650.00.xlsx"

# 价格唯一来源：结算清单。明细行与合计都以它为准。
SETTLE_XLSX = DOC_DIR / "9、结算清单 价格.xlsx"

# 红色系：覆盖 EE0000 / FF0000
RED_R_MIN, RED_G_MAX, RED_B_MAX = 0x80, 0x60, 0x60


def is_red(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    c = rpr.find(qn("w:color"))
    if c is None:
        return False
    val = c.get(qn("w:val"))
    if not val or len(val) != 6:
        return False
    try:
        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except ValueError:
        return False
    return r >= RED_R_MIN and g <= RED_G_MAX and b <= RED_B_MAX


def red_blocks(para):
    """把相邻的红色 run 合并成逻辑块——Word 常把一个词切成多个 run。"""
    out, buf = [], []
    for r in para.runs:
        if is_red(r):
            buf.append(r)
        elif buf:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out


def all_paragraphs(doc):
    """按文档顺序产出所有段落：正文 → 表格单元格 → 页眉页脚。"""
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                k = cell._tc.getroottree().getpath(cell._tc)
                if k in seen:
                    continue
                seen.add(k)
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def uniq_cells(row):
    """合并单元格会被重复返回，按 tc 的 XML 路径去重。"""
    out, seen = [], set()
    for c in row.cells:
        k = c._tc.getroottree().getpath(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def set_cell(cell, text):
    """写值并改黑：只改第一个 run，其余清空，保留字体等格式。"""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for r in p.runs[1:]:
            r.text = ""
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0, 0, 0)


def fmt_val(v, money=False):
    if v is None:
        return ""
    if money:
        try:
            return f"{float(v):.2f}"
        except (TypeError, ValueError):
            return str(v).strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _load_info(xlsx):
    if not xlsx.exists():
        raise FileNotFoundError(f"找不到红字信息：{xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    by_ph, by_field = {}, {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 3:
            continue
        ph, field, val = row[0], row[1], row[2]
        if val is None:
            continue
        text = fmt_val(val, "金额" in str(field) and "大写" not in str(field))
        if ph and str(ph).strip() not in ("—", ""):
            by_ph[str(ph).strip()] = text
        if field:
            by_field[str(field).strip()] = text
    wb.close()
    return by_ph, by_field


def load_red_info(company):
    """读某公司的红字信息表，返回 (按占位取值, 按字段名取值)。"""
    return _load_info(INFO_DIR / f"{company}.xlsx")


_OWNER_OVERRIDE = {}


def set_owner_override(name=None, project=None, short=None, address=None):
    """GUI 里手填的甲方信息，覆盖 甲方.xlsx 里的值。传 None 的项保持原样。"""
    global _OWNER_OVERRIDE
    ov = dict(_OWNER_OVERRIDE)
    if name:
        ov["某某甲方"] = name
        ov["采购人名称"] = name          # 签收单的占位写法
        ov["某某甲方内"] = f"{name}校内"
    if project:
        ov["某某甲方某某项目"] = project
        ov["某某甲方项目名字"] = project
    if short:
        ov["某某文化墙"] = short
    elif project:
        # 没给简称就从项目名里剥掉甲方名称：「宁夏小学开放式书吧」→「开放式书吧」
        tail = project[len(name):].strip() if name and project.startswith(name) else ""
        ov["某某文化墙"] = tail or project
    if address:
        ov["某某甲方地址"] = address
    elif name:
        # 换了甲方却没给地址，就不能沿用旧项目的地址，填「空」占位
        ov["某某甲方地址"] = "空"
    _OWNER_OVERRIDE = ov


def reset_owner_override():
    global _OWNER_OVERRIDE
    _OWNER_OVERRIDE = {}


def load_owner_info():
    """读甲方红字信息表（红字信息/甲方.xlsx），GUI 手填的会覆盖同名占位。"""
    by_ph, by_field = _load_info(INFO_DIR / "甲方.xlsx")
    if _OWNER_OVERRIDE:
        old_name = by_ph.get("某某甲方")
        new_name = _OWNER_OVERRIDE.get("某某甲方")
        if old_name and new_name and old_name != new_name:
            # 模板里直接写死旧甲方真名的红字（如签收单的「采购人名称」），一并换成新甲方
            for k, v in list(by_ph.items()):
                if v == old_name:
                    by_ph[k] = new_name
        by_ph.update(_OWNER_OVERRIDE)
    return by_ph, by_field


def load_settlement():
    """读结算清单：返回 (明细行, 合计金额, 合计大写)。价格以此为准。"""
    from sync_quote_table import read_quote
    return read_quote(SETTLE_XLSX)


MONEY_NUM = re.compile(r"[\d,]+(\.\d+)?")
MONEY_CN = re.compile(r"[零壹贰叁肆伍陆柒捌玖拾佰仟万亿元角分整]{3,}")


def apply_settlement(by_ph, amt, cn):
    """把映射里所有金额类占位（数字或中文大写）的值改指向结算清单。

    占位符文本保持原样——它只是用来匹配模板里红字的；值换成结算清单的金额。
    """
    out = dict(by_ph)
    for k in out:
        key = str(k).strip()
        if MONEY_NUM.fullmatch(key):
            out[k] = f"{amt:.2f}"
        elif MONEY_CN.fullmatch(key):
            out[k] = cn
    return out


def replace_in_doc(doc, by_ph, resolver=None):
    """替换文档里的红字块。

    resolver(para_text, red_text, by_ph) 可返回覆盖值或 None；
    不传 resolver 时直接查 by_ph。
    返回 (命中列表, 未命中列表)。
    """
    hits, miss = [], []
    for para in all_paragraphs(doc):
        ptext = para.text
        for buf in red_blocks(para):
            raw = "".join(r.text for r in buf)
            text = raw.strip()
            if not text:
                continue
            new = resolver(ptext, text, by_ph) if resolver else by_ph.get(text)
            if new is None:
                miss.append(text)
                continue
            # 红字块可能带着尾随空格（排版间隔），替换后要还回去，
            # 否则「某某甲方␣␣」+「某某公司」会粘成「宁夏小学某某公司」
            trailing = raw[len(raw.rstrip()):]
            buf[0].text = new
            buf[0].font.color.rgb = RGBColor(0, 0, 0)
            for r in buf[1:]:
                r.text = ""
            if trailing and len(buf) > 1:
                buf[-1].text = trailing
                buf[-1].font.color.rgb = RGBColor(0, 0, 0)
            hits.append((text, new))
    return hits, miss


def merge_maps(*maps):
    out = {}
    for m in maps:
        out.update(m)
    return out


def strip_ws(s):
    return re.sub(r"\s+", "", str(s))
