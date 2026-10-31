# -*- coding: utf-8 -*-
"""通用红字填充：用指定公司的红字信息替换 docx 里的红色系标记。

用法：
    python tools/fill_docx.py <源docx> <输出docx> [公司名，默认晨早]

只替换红字信息表里存在的占位，其余（甲方、表格明细等）保持原样。
"""
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:/git学习/Excel转指定位置word")
INFO_DIR = ROOT / "项目文件" / "红字信息"


def is_red(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    c = rpr.find(qn("w:color"))
    if c is None:
        return False
    val = c.get(qn("w:val"))
    if not val or len(val) != 6:
        return False
    try:
        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except ValueError:
        return False
    return r >= 0x80 and g <= 0x60 and b <= 0x60


def red_blocks(para):
    out, buf = [], []
    for r in para.runs:
        if is_red(r):
            buf.append(r)
        elif buf:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                k = cell._tc.getroottree().getpath(cell._tc)
                if k in seen:
                    continue
                seen.add(k)
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def fmt(val, field=""):
    if isinstance(val, (int, float)) and "金额" in str(field) and "大写" not in str(field):
        return f"{float(val):.2f}"
    if isinstance(val, float):
        return f"{val:.2f}"
    return str(val).strip()


def load_mapping(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    mapping = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 3:
            continue
        ph, field, val = row[0], row[1], row[2]
        if not ph or str(ph).strip() in ("—", "") or val is None:
            continue
        mapping[str(ph).strip()] = fmt(val, field)
    wb.close()
    return mapping


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    company = sys.argv[3] if len(sys.argv) > 3 else "晨早"

    xlsx = INFO_DIR / f"{company}.xlsx"
    if not xlsx.exists():
        sys.exit(f"找不到红字信息：{xlsx}")

    merged = load_mapping(xlsx)
    print(f"载入 {company}.xlsx：{len(merged)} 条")

    doc = Document(str(src))
    hits, miss = [], []
    for para in all_paragraphs(doc):
        for buf in red_blocks(para):
            text = "".join(r.text for r in buf).strip()
            if not text:
                continue
            if text in merged:
                buf[0].text = merged[text]
                buf[0].font.color.rgb = RGBColor(0, 0, 0)
                for r in buf[1:]:
                    r.text = ""
                hits.append((text, merged[text]))
            else:
                miss.append(text)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    print("\n" + "=" * 72)
    print(f"已替换 {len(hits)} 处：")
    for old, new in hits:
        print(f"  {old}  →  {new}")
    if miss:
        print(f"\n未替换 {len(miss)} 处（红字信息中无此项，保持红字）：")
        for t, n in Counter(miss).most_common():
            label = t if len(t) <= 28 else t[:28] + "…"
            print(f"  {label}  ×{n}")
    print("\n输出：", out)


if __name__ == "__main__":
    main()
