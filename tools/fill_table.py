# -*- coding: utf-8 -*-
"""把晨早报价单的明细表原样搬进 合同 / 签收单 / 验收书 的表格里（覆盖 _已填 产物，不动源文件）。"""
import sys
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import RGBColor

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:/git学习/Excel转指定位置word")
BASE = ROOT / "项目文件" / "1、工程A三方询价 价格"
QUOTE = BASE / "3.报价" / "1、报价清单（晨早）44650.00.xlsx"

TARGETS = [
    (BASE / "6、合同_已填.docx", "合同"),
    (BASE / "7、签收单_已填.docx", "签收单"),
    (BASE / "8、验收书_已填.docx", "验收书"),
]


def read_quote():
    """返回 8 行明细：[(序号, 名称, 参数, 数量, 单位, 单价, 总价), ...]"""
    wb = openpyxl.load_workbook(QUOTE, data_only=True)
    ws = wb.active
    rows = []
    for r in range(3, 11):          # 报价单第 3~10 行是 8 条明细
        vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
        rows.append(vals)
    wb.close()
    return rows


def uniq_cells(row):
    out, seen = [], set()
    for c in row.cells:
        k = c._tc.getroottree().getpath(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def find_data_start(table):
    """定位第一条明细所在行（第 0 列是 1、第 1 列是人物造型）。"""
    for i, row in enumerate(table.rows):
        cells = uniq_cells(row)
        if len(cells) >= 3 and cells[0].text.strip() == "1" and "人物" in cells[1].text:
            return i
    return None


def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for r in p.runs[1:]:
            r.text = ""
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0, 0, 0)


def fmt(v, is_money=False):
    if v is None:
        return ""
    if is_money:
        try:
            return f"{float(v):.2f}"
        except (TypeError, ValueError):
            return str(v).strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


quote = read_quote()
print(f"报价单明细 {len(quote)} 行：")
for q in quote:
    print("   ", " | ".join(fmt(v, i in (5, 6)) for i, v in enumerate(q)))

for path, label in TARGETS:
    if not path.exists():
        print(f"\n[跳过] {label}：文件不存在 {path}")
        continue
    doc = Document(str(path))
    table = doc.tables[0]
    start = find_data_start(table)
    if start is None:
        print(f"\n[跳过] {label}：没定位到明细起始行")
        continue

    changed = 0
    for k, vals in enumerate(quote):
        cells = uniq_cells(table.rows[start + k])
        for ci in range(min(7, len(cells))):
            new = fmt(vals[ci], ci in (5, 6))
            if cells[ci].text.strip() != new:
                changed += 1
            set_cell(cells[ci], new)

    doc.save(str(path))
    print(f"\n[完成] {label}：起始行 R{start}，写入 {len(quote)} 行，文本变更 {changed} 格")
