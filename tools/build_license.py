# -*- coding: utf-8 -*-
"""为三家公司生成「营业执照」页：一张 A4，只有标题和一个粘贴框。

真实营业执照扫描件需要企业提供，脚本只负责排版占位。
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"D:/git学习/Excel转指定位置word/项目文件/1、工程A三方询价 价格/2.资质")
FONT = "宋体"

COMPANIES = [
    ("1、晨早", "长沙晨早文化传播有限公司"),
    ("2、泉旭", "长沙泉旭广告传媒有限公司"),
    ("3、左右", "长沙左右文化传播有限公司"),
]


def set_font(run, size=12, bold=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def build(folder, company):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.5)

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(company), 16, True, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("营业执照（复印件）")
    set_font(r, 12)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # 粘贴框：1×1 表格，撑满版面
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Cm(16)
    table.rows[0].height = Cm(18)

    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].paragraph_format.space_before = Pt(200)
    r = cell.paragraphs[0].add_run("此处粘贴营业执照扫描件")
    set_font(r, 12)
    r.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("（或加盖公章的复印件，扫描后插入图片）")
    set_font(r2, 10)
    r2.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)

    out = BASE / folder / "营业执照.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


for folder, company in COMPANIES:
    print("saved:", build(folder, company))
