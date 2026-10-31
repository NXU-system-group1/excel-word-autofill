# -*- coding: utf-8 -*-
"""用「红字信息」Excel 的值替换声明 docx 里的红色系标记，输出到新目录，不动源文件。"""
import shutil
import sys
from pathlib import Path

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"D:/git学习/Excel转指定位置word/项目文件")
INFO_DIR = BASE / "红字信息"
SRC_DIR = BASE / "1、工程A三方询价 价格" / "1.声明"
OUT_DIR = SRC_DIR / "已填"

# 文件名关键词 -> 红字信息 Excel
MATCH = {
    "晨早": "晨早.xlsx",
    "泉旭": "泉旭.xlsx",
    "左右": "左右.xlsx",
}


def is_red(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    c = rpr.find(qn("w:color"))
    if c is None:
        return False
    val = c.get(qn("w:val"))
    if not val or len(val) != 6:
        return False
    try:
        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except ValueError:
        return False
    return r >= 0x80 and g <= 0x60 and b <= 0x60


def red_blocks(para):
    out, buf = [], []
    for r in para.runs:
        if is_red(r):
            buf.append(r)
        elif buf:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                k = cell._tc.getroottree().getpath(cell._tc)
                if k in seen:
                    continue
                seen.add(k)
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def load_mapping(xlsx_path):
    """模板红字占位 -> 值。占位为「—」或空的行跳过。"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    mapping = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 3:
            continue
        ph, field, val = row[0], row[1], row[2]
        if not ph or ph == "—" or val is None:
            continue
        mapping[str(ph).strip()] = str(val).strip()
    wb.close()
    return mapping


OUT_DIR.mkdir(parents=True, exist_ok=True)

print("=" * 76)
for src in sorted(SRC_DIR.glob("*.docx")):
    if src.name.startswith("~$"):
        continue
    key = next((k for k in MATCH if k in src.name), None)
    dst = OUT_DIR / src.name

    if key is None:
        shutil.copy2(src, dst)
        print(f"[跳过] {src.name}：无对应红字信息，原样复制")
        continue

    mapping = load_mapping(INFO_DIR / MATCH[key])
    doc = Document(str(src))
    hits, miss = [], []

    for para in all_paragraphs(doc):
        for buf in red_blocks(para):
            text = "".join(r.text for r in buf).strip()
            if not text:
                continue
            if text in mapping:
                buf[0].text = mapping[text]
                buf[0].font.color.rgb = RGBColor(0, 0, 0)
                for r in buf[1:]:
                    r.text = ""
                hits.append((text, mapping[text]))
            else:
                miss.append(text)

    doc.save(str(dst))
    print(f"[填充] {src.name}  ←  {MATCH[key]}")
    for old, new in hits:
        print(f"        {old}  →  {new}")
    if miss:
        print(f"        未替换（红字信息中无此项）：{'、'.join(sorted(set(miss)))}")

print("=" * 76)
print("输出目录：", OUT_DIR)
